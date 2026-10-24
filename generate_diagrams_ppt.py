import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_diagrams_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    NAVY_PRIMARY = RGBColor(26, 54, 93)
    BG_DARK = RGBColor(15, 23, 42)
    BLUE_ACCENT = RGBColor(56, 189, 248)
    TEXT_LIGHT = RGBColor(248, 250, 252)
    TEXT_MUTED = RGBColor(148, 163, 184)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide for Section 4 Diagrams
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    bg.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(3.0))
    tf = tb.text_frame
    p1 = tf.paragraphs[0]
    p1.text = "SECTION 4: SYSTEM DESIGN & UML MODELING"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT
    p1.font.name = 'Segoe UI'

    p2 = tf.add_paragraph()
    p2.text = "Smart College Analytics System (SCAS) - Complete Architectural & UML Diagrams"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = BLUE_ACCENT
    p2.font.name = 'Segoe UI'

    p3 = tf.add_paragraph()
    p3.text = "Gantt Chart | E-R Diagram | Use Case | Activity | Sequence | Collaboration | Class | 3-Tier Architecture"
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED
    p3.font.name = 'Segoe UI'

    # -------------------------------------------------------------
    # DIAGRAM SLIDES (Full-Slide Visual Pictures)
    # -------------------------------------------------------------
    diagram_files = [
        ("scas_timeline_chart.png", "4.1 Timeline Chart (Gantt Chart)"),
        ("scas_er_diagram.png", "4.3 Entity-Relationship (E-R) Diagram"),
        ("scas_usecase_diagram.png", "4.4.1 UML Use Case Diagram"),
        ("scas_activity_diagram.png", "4.4.2 UML Activity Diagram"),
        ("scas_sequence_diagram.png", "4.4.3 UML Sequence Diagram"),
        ("scas_collaboration_diagram.png", "4.4.4 UML Collaboration Diagram"),
        ("scas_class_diagram.png", "4.4.5 UML Class Diagram"),
        ("scas_system_architecture_dfd.png", "3-Tier System Architecture & DFD")
    ]

    for filename, title in diagram_files:
        img_path = os.path.join(os.getcwd(), filename)
        if os.path.exists(img_path):
            s = prs.slides.add_slide(blank_layout)
            # Add full-screen picture shape
            s.shapes.add_picture(img_path, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))
            print(f"Added slide: {title}")

    out_path = os.path.join(os.getcwd(), "SCAS_Project_Diagrams.pptx")
    prs.save(out_path)
    print(f"\nPresentation successfully created at: {out_path}")

if __name__ == "__main__":
    create_diagrams_presentation()
