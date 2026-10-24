import os
import sys
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # Set slide dimensions to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    CARD_BG = RGBColor(30, 41, 59)       # Slate 800
    NAVY_PRIMARY = RGBColor(26, 54, 93)   # #1A365D
    BLUE_ACCENT = RGBColor(49, 130, 206) # #3182CE
    TEXT_LIGHT = RGBColor(248, 250, 252)  # White/Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400
    TEXT_DARK = RGBColor(30, 41, 59)

    blank_slide_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="SCAS - PROJECT PRESENTATION"):
        # Header background banner
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        banner.fill.solid()
        banner.fill.fore_color.rgb = NAVY_PRIMARY
        banner.line.fill.background()

        # Category text
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.3))
        tf = txBox.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = BLUE_ACCENT
        p.font.name = 'Segoe UI'

        # Main Title text
        txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(11.7), Inches(0.55))
        tf2 = txBox2.text_frame
        tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_LIGHT
        p2.font.name = 'Segoe UI'

        # Footer
        tx_foot = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(11.7), Inches(0.3))
        p_f = tx_foot.text_frame.paragraphs[0]
        p_f.text = "Smart College Analytics System (SCAS) | Department of Computer Science & Applications"
        p_f.font.size = Pt(9)
        p_f.font.color.rgb = TEXT_MUTED
        p_f.font.name = 'Segoe UI'

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_DARK
    bg1.line.fill.background()

    # Title box
    tb = slide1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(2.0))
    p = tb.text_frame.paragraphs[0]
    p.text = "SMART COLLEGE ANALYTICS SYSTEM"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT
    p.font.name = 'Segoe UI'

    p_sub = tb.text_frame.add_paragraph()
    p_sub.text = "(SCAS) - MCA Continuous Assessment & Academic Analytics Platform"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = BLUE_ACCENT
    p_sub.font.name = 'Segoe UI'

    # Divider line
    line = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(4.3), Inches(11.333), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE_ACCENT
    line.line.fill.background()

    # Presenter Details
    tb_meta = slide1.shapes.add_textbox(Inches(1.0), Inches(4.7), Inches(11.333), Inches(2.0))
    p_m1 = tb_meta.text_frame.paragraphs[0]
    p_m1.text = "A Project Presentation Submitted for Master of Computer Applications (MCA)"
    p_m1.font.size = Pt(14)
    p_m1.font.color.rgb = TEXT_MUTED
    p_m1.font.name = 'Segoe UI'

    p_m2 = tb_meta.text_frame.add_paragraph()
    p_m2.text = "Department of Computer Science & Applications"
    p_m2.font.size = Pt(14)
    p_m2.font.bold = True
    p_m2.font.color.rgb = TEXT_LIGHT
    p_m2.font.name = 'Segoe UI'

    # -------------------------------------------------------------
    # SLIDE 2: 1. INTRODUCTION
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "1. INTRODUCTION & OVERVIEW", "1. INTRODUCTION")

    box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3))
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "About Smart College Analytics System (SCAS):"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    bullets = [
        "The Smart College Analytics System (SCAS) is an enterprise-grade higher education ERP (Enterprise Resource Planning) and continuous assessment analytics portal.",
        "It provides a multi-tenant, role-based ecosystem designed to streamline academic administrative operations for Administrators, Faculty Members, and Students.",
        "SCAS bridges manual attendance registers, internal mark calculations, fee collection tracking, leave requests, and placement readiness into a unified digital dashboard.",
        "Goal of the System: To digitalize academic tracking, eliminate administrative delays, provide real-time performance analytics, and enhance overall institutional efficiency."
    ]

    for b in bullets:
        p_b = tf.add_paragraph()
        p_b.text = "• " + b
        p_b.font.size = Pt(14)
        p_b.font.color.rgb = TEXT_DARK
        p_b.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 3: 1.1 Project Profile & 1.2 About Organization
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "1.1 Project Profile & 1.2 About the Organization", "1. INTRODUCTION")

    # Left Card (Project Profile)
    card1 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
    card1.fill.solid()
    card1.fill.fore_color.rgb = RGBColor(241, 245, 249)
    card1.line.color.rgb = BLUE_ACCENT

    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = Inches(0.3)
    p = tf1.paragraphs[0]
    p.text = "1.1 Project Profile"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    pp_items = [
        ("Project Title: ", "Smart College Analytics System (SCAS)"),
        ("Project Type: ", "Web-based Education ERP & Analytics Portal"),
        ("Course / Level: ", "Master of Computer Applications (MCA - PS3)"),
        ("Front-End Tech: ", "React 19, TypeScript, Vite, Tailwind CSS"),
        ("Back-End Tech: ", "Python 3.x, Flask Framework, SQLAlchemy ORM"),
        ("Database: ", "MySQL / SQLite Relational Database"),
        ("Primary Users: ", "Administrator, Faculty, and Students")
    ]
    for label, val in pp_items:
        p_i = tf1.add_paragraph()
        r1 = p_i.add_run()
        r1.text = "• " + label
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = NAVY_PRIMARY
        r2 = p_i.add_run()
        r2.text = val
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_DARK
        p_i.space_after = Pt(6)

    # Right Card (About Organization)
    card2 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
    card2.fill.solid()
    card2.fill.fore_color.rgb = RGBColor(241, 245, 249)
    card2.line.color.rgb = BLUE_ACCENT

    tf2 = card2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = Inches(0.3)
    p2 = tf2.paragraphs[0]
    p2.text = "1.2 About the Organization"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = NAVY_PRIMARY

    org_items = [
        "Academic Institution: Department of Computer Science & Applications / Higher Education College.",
        "Focus Area: Delivering quality postgraduate computer education, continuous assessment, and research excellence.",
        "Digital Transformation: Transitioning from manual paper-based record-keeping to a automated, cloud-ready digital infrastructure.",
        "Commitment: Enhancing student career readiness through placement aptitude tracking, transparent fee monitoring, and real-time faculty-student interaction."
    ]
    for o in org_items:
        p_o = tf2.add_paragraph()
        p_o.text = "• " + o
        p_o.font.size = Pt(12)
        p_o.font.color.rgb = TEXT_DARK
        p_o.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 4: 2. SYSTEM STUDY AND ANALYSIS - Problem & Existing System
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "2.1 Problem Statement & 2.2 Existing System", "2. SYSTEM STUDY AND ANALYSIS")

    # Card 1: Problem Statement
    c1 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
    c1.fill.solid()
    c1.fill.fore_color.rgb = RGBColor(254, 242, 242) # Light red tint
    c1.line.color.rgb = RGBColor(239, 68, 68)

    t1 = c1.text_frame
    t1.word_wrap = True
    t1.margin_left = t1.margin_top = t1.margin_right = t1.margin_bottom = Inches(0.3)
    p = t1.paragraphs[0]
    p.text = "2.1 Problem Statement"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(185, 28, 28)

    probs = [
        "Academic institutions face severe inefficiencies due to fragmented, paper-based administrative operations.",
        "Manual calculation of internal marks, midterms, and GPAs leads to high error rates and delayed grade sheets.",
        "Lack of real-time attendance monitoring results in undetected attendance shortages (<75%) until end-semester.",
        "Student fee management relies on physical ledgers and offline counters, causing long queues and payment tracking delays.",
        "Leave and duty applications rely on physical paper forms, leading to lost requests and delayed approvals."
    ]
    for pr in probs:
        p_p = t1.add_paragraph()
        p_p.text = "• " + pr
        p_p.font.size = Pt(12)
        p_p.font.color.rgb = TEXT_DARK
        p_p.space_after = Pt(6)

    # Card 2: Existing System & Drawbacks
    c2 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
    c2.fill.solid()
    c2.fill.fore_color.rgb = RGBColor(254, 242, 242)
    c2.line.color.rgb = RGBColor(239, 68, 68)

    t2 = c2.text_frame
    t2.word_wrap = True
    t2.margin_left = t2.margin_top = t2.margin_right = t2.margin_bottom = Inches(0.3)
    p2 = t2.paragraphs[0]
    p2.text = "2.2 Existing System & 2.2.1 Drawbacks"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(185, 28, 28)

    drawbacks = [
        "Existing Approach: Paper registers, isolated Excel spreadsheets, physical notice boards, and separate offline fee counters.",
        "Drawback 1: Data Redundancy & Inconsistency across departments.",
        "Drawback 2: High Administrative Overhead and human calculation errors.",
        "Drawback 3: No Centralized Analytics or visual graphs for student progress tracking.",
        "Drawback 4: Delayed Communication for notices, password resets, and fee dues.",
        "Drawback 5: Lack of Career Preparation integration like mock aptitude tests."
    ]
    for d in drawbacks:
        p_d = t2.add_paragraph()
        p_d.text = "• " + d
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_DARK
        p_d.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 5: 2.3 Proposed System & Comparison Study
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "2.3 Proposed System & 2.3.1 Comparison Study", "2. SYSTEM STUDY AND ANALYSIS")

    # Table comparing Existing vs Proposed
    rows, cols = 7, 3
    left, top, width, height = Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2)
    table_shape = slide5.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.7)
    table.columns[1].width = Inches(4.5)
    table.columns[2].width = Inches(4.5)

    headers = ["Feature / Aspect", "Existing Manual System", "Proposed SCAS Platform"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY_PRIMARY
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_LIGHT

    comp_data = [
        ["Attendance Tracking", "Paper registers, manual tallying at end of semester", "Real-time session marking with automated low-attendance warnings (<75%)"],
        ["Mark & GPA Processing", "Manual Excel sheets, error-prone grade calculation", "Automatic SGPA/CGPA computation with dynamic Chart.js analytics graphs"],
        ["Notice Board & Alerts", "Physical paper flyers pinned on college walls", "Digital Notice Board with instant pinned notices and SMTP email alerts"],
        ["Leave & Duty Requests", "Physical written letters submitted manually", "Multi-tier online leave application submission and real-time faculty approval"],
        ["Fee & Payment Dues", "Manual ledger books and separate offline counter", "Real-time fee status tracking (Paid/Pending) with auto-generated receipt numbers"],
        ["Placement & Aptitude", "No integrated system for placement preparation", "Built-in mock placement aptitude test engine with instant scoreboards"]
    ]

    for row_idx, row in enumerate(comp_data):
        for col_idx, cell_value in enumerate(row):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = cell_value
            cell.fill.solid()
            bg_col = RGBColor(241, 245, 249) if row_idx % 2 == 1 else RGBColor(255, 255, 255)
            cell.fill.fore_color.rgb = bg_col
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_DARK
            if col_idx == 0:
                p.font.bold = True
                p.font.color.rgb = NAVY_PRIMARY

    # -------------------------------------------------------------
    # SLIDE 6: 2.3.2 System Modules
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "2.3.2 System Modules Breakdown", "2. SYSTEM STUDY AND ANALYSIS")

    modules = [
        ("1. Auth & Security Module: ", "Role-Based Access Control (RBAC) with password hashing (Werkzeug) and session handling for Admin, Faculty, and Students."),
        ("2. Admin Management Module: ", "Central dashboard for managing student/faculty profiles, generating credentials, publishing notices, and fee tracking."),
        ("3. Faculty Management Module: ", "Subject allocation, real-time daily attendance marker, internal marks upload engine, and leave application approval workflow."),
        ("4. Student Portal Module: ", "Academic dashboard, SGPA/CGPA grade sheets, subject attendance logs, performance analytics charts, and leave requests."),
        ("5. Fee & Payment Module: ", "Tracks tuition payments, updates payment status ('Paid'/'Pending'), and generates unique transaction receipts."),
        ("6. Placement & Aptitude Module: ", "Mock aptitude exam engine allowing faculty to create tests and students to practice for placement drives."),
        ("7. Email & Notification Engine: ", "SMTP integration sending welcome credentials, temporary password resets, and alert notifications.")
    ]

    box = slide6.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(5.5))
    tf = box.text_frame
    tf.word_wrap = True

    for title, desc in modules:
        p = tf.add_paragraph()
        r1 = p.add_run()
        r1.text = title
        r1.font.bold = True
        r1.font.size = Pt(13)
        r1.font.color.rgb = NAVY_PRIMARY

        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_DARK
        p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 7: 3. DEVELOPMENT ENVIRONMENT - Requirements
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "3.1 Hardware & 3.2 Software Requirements", "3. DEVELOPMENT ENVIRONMENT")

    # Left Box: Hardware
    card_h = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
    card_h.fill.solid()
    card_h.fill.fore_color.rgb = RGBColor(241, 245, 249)
    card_h.line.color.rgb = BLUE_ACCENT

    tf_h = card_h.text_frame
    tf_h.word_wrap = True
    tf_h.margin_left = tf_h.margin_top = tf_h.margin_right = tf_h.margin_bottom = Inches(0.3)

    p = tf_h.paragraphs[0]
    p.text = "3.1 Hardware Requirements"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    hw_specs = [
        ("Processor: ", "Intel Core i3 / i5 / AMD Ryzen 5 (2.0 GHz or higher)"),
        ("RAM: ", "Minimum 4 GB (8 GB recommended for optimal performance)"),
        ("Storage / Hard Disk: ", "Minimum 20 GB available SSD / HDD storage"),
        ("Display: ", "1366 x 768 minimum resolution (1920 x 1080 recommended)"),
        ("Network: ", "Standard LAN / Wi-Fi internet connectivity for SMTP & client access")
    ]
    for label, val in hw_specs:
        p_i = tf_h.add_paragraph()
        r1 = p_i.add_run()
        r1.text = "• " + label
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = NAVY_PRIMARY
        r2 = p_i.add_run()
        r2.text = val
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_DARK
        p_i.space_after = Pt(10)

    # Right Box: Software
    card_s = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
    card_s.fill.solid()
    card_s.fill.fore_color.rgb = RGBColor(241, 245, 249)
    card_s.line.color.rgb = BLUE_ACCENT

    tf_s = card_s.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = Inches(0.3)

    p2 = tf_s.paragraphs[0]
    p2.text = "3.2 Software Requirements"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = NAVY_PRIMARY

    sw_specs = [
        ("Operating System: ", "Windows 10 / 11, Linux (Ubuntu/Debian), or macOS"),
        ("Web Server: ", "Flask WSGI Server / Gunicorn (Production)"),
        ("Runtime Environments: ", "Python 3.10+ and Node.js v18+"),
        ("Database Server: ", "MySQL Server 8.0 (via XAMPP/WAMP) or SQLite3"),
        ("Development Tools: ", "VS Code, Git, Postman, Vite CLI Engine"),
        ("Web Browser: ", "Google Chrome, Microsoft Edge, Mozilla Firefox")
    ]
    for label, val in sw_specs:
        p_i = tf_s.add_paragraph()
        r1 = p_i.add_run()
        r1.text = "• " + label
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = NAVY_PRIMARY
        r2 = p_i.add_run()
        r2.text = val
        r2.font.size = Pt(12)
        r2.font.color.rgb = TEXT_DARK
        p_i.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 8: 3.3 Programming Environment - Front-End & Back-End
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "3.3 Programming Environment (Front-End & Back-End)", "3. DEVELOPMENT ENVIRONMENT")

    # Left: Front-End
    card_fe = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
    card_fe.fill.solid()
    card_fe.fill.fore_color.rgb = RGBColor(241, 245, 249)
    card_fe.line.color.rgb = BLUE_ACCENT

    tf_fe = card_fe.text_frame
    tf_fe.word_wrap = True
    tf_fe.margin_left = tf_fe.margin_top = tf_fe.margin_right = tf_fe.margin_bottom = Inches(0.3)

    p = tf_fe.paragraphs[0]
    p.text = "3.3.1 About Front-End"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    fe_techs = [
        ("React 19 & TypeScript: ", "Component-based Single Page Application (SPA) architecture offering strict type safety and modular components."),
        ("Vite Build Engine: ", "Blazing fast Hot Module Replacement (HMR) development server and optimized bundle generator."),
        ("Tailwind CSS: ", "Utility-first CSS framework enabling dark mode glassmorphism, responsive grid layouts, and custom animations."),
        ("Chart.js & react-chartjs-2: ", "Dynamic interactive charts displaying GPA progression, attendance distribution, and subject breakdowns."),
        ("Lucide React Icons: ", "Modern vector iconography providing clean UI controls across all portal views.")
    ]
    for label, val in fe_techs:
        p_i = tf_fe.add_paragraph()
        r1 = p_i.add_run()
        r1.text = "• " + label
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = NAVY_PRIMARY
        r2 = p_i.add_run()
        r2.text = val
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_DARK
        p_i.space_after = Pt(6)

    # Right: Back-End
    card_be = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
    card_be.fill.solid()
    card_be.fill.fore_color.rgb = RGBColor(241, 245, 249)
    card_be.line.color.rgb = BLUE_ACCENT

    tf_be = card_be.text_frame
    tf_be.word_wrap = True
    tf_be.margin_left = tf_be.margin_top = tf_be.margin_right = tf_be.margin_bottom = Inches(0.3)

    p2 = tf_be.paragraphs[0]
    p2.text = "3.3.2 About Back-End"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = NAVY_PRIMARY

    be_techs = [
        ("Python 3.x & Flask: ", "Lightweight, extensible WSGI web server framework handling REST API routes and blueprint controllers."),
        ("Flask-SQLAlchemy (ORM): ", "Object-Relational Mapping library translating Python model classes into MySQL/SQLite SQL queries."),
        ("MySQL Database: ", "Relational database storing 12 normalized tables with primary keys, foreign key constraints, and cascade rules."),
        ("Werkzeug Security: ", "Secure password hashing (generate_password_hash) and checking for safe user authentication."),
        ("SMTP Mail Engine: ", "Automated email transport service for dispatching default login credentials and password reset notifications.")
    ]
    for label, val in be_techs:
        p_i = tf_be.add_paragraph()
        r1 = p_i.add_run()
        r1.text = "• " + label
        r1.font.bold = True
        r1.font.size = Pt(12)
        r1.font.color.rgb = NAVY_PRIMARY
        r2 = p_i.add_run()
        r2.text = val
        r2.font.size = Pt(11)
        r2.font.color.rgb = TEXT_DARK
        p_i.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 9: Conclusion & Summary
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide9, "SUMMARY & CONCLUSION", "PROJECT SUMMARY")

    box = slide9.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3))
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Project Summary & Outcomes:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY_PRIMARY

    summary_bullets = [
        "SCAS successfully replaces error-prone manual paper processes with an automated, multi-tenant digital ERP portal.",
        "Real-time attendance marking and automated risk alerts (<75%) empower faculty and students to take timely action.",
        "Automatic GPA/SGPA calculation combined with Chart.js analytics provides instant insights into academic performance.",
        "Integrated fee tracking and receipt generation streamlines financial record-keeping for administrative staff.",
        "Built-in placement aptitude testing helps students prepare for campus recruitment drives effectively.",
        "System diagnostic routines (verify_system.py) and database seeders ensure high reliability and easy deployment."
    ]

    for sb in summary_bullets:
        p_b = tf.add_paragraph()
        p_b.text = "✔ " + sb
        p_b.font.size = Pt(14)
        p_b.font.color.rgb = TEXT_DARK
        p_b.space_after = Pt(10)

    # Save presentation
    ppt_path = os.path.join(os.getcwd(), "SCAS_Project_Presentation.pptx")
    prs.save(ppt_path)
    print(f"Presentation successfully created at: {ppt_path}")

if __name__ == "__main__":
    create_presentation()
