import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_info_document():
    doc = docx.Document()

    # Define color palette
    PRIMARY_COLOR = RGBColor(26, 54, 93)     # Deep Navy (#1A365D)
    SECONDARY_COLOR = RGBColor(49, 130, 206) # Blue Accent (#3182CE)
    DARK_TEXT = RGBColor(45, 55, 72)         # Slate Gray (#2D3748)
    MUTED_TEXT = RGBColor(113, 128, 150)     # Muted Gray (#718096)

    # Set Margins (1 inch all sides)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = DARK_TEXT

    # Helper Functions
    def add_main_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(20)
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.italic = True
        run.font.color.rgb = SECONDARY_COLOR

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = SECONDARY_COLOR

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR

    def add_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(11)
        run.font.color.rgb = DARK_TEXT
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.font.name = 'Segoe UI'
            run_b.font.bold = True
            run_b.font.color.rgb = PRIMARY_COLOR
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.color.rgb = DARK_TEXT

    def add_callout(text, title="NOTE"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        
        # Set shading background
        shading = parse_xml(r'<w:shd {} w:fill="EBF8FF"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        
        # Set left border thick blue
        borders = parse_xml(r'''
            <w:tcBorders {} >
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="36" w:space="0" w:color="3182CE"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        '''.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.1)
        
        r_t = p.add_run(f"📌 {title}: ")
        r_t.bold = True
        r_t.font.color.rgb = SECONDARY_COLOR
        
        r_b = p.add_run(text)
        r_b.font.color.rgb = DARK_TEXT
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def style_table(table, col_widths, headers, data):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            hdr_cells[i].text = header
            hdr_cells[i].width = Inches(col_widths[i])
            shd = parse_xml(r'<w:shd {} w:fill="1A365D"/>'.format(nsdecls('w')))
            hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = 'Segoe UI'
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)

        for row_idx, row_data in enumerate(data):
            row_cells = table.add_row().cells
            bg_color = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, cell_value in enumerate(row_data):
                row_cells[col_idx].text = str(cell_value)
                row_cells[col_idx].width = Inches(col_widths[col_idx])
                shd = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg_color))
                row_cells[col_idx]._tc.get_or_add_tcPr().append(shd)
                p = row_cells[col_idx].paragraphs[0]
                for run in p.runs:
                    run.font.name = 'Segoe UI'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = DARK_TEXT

    # -------------------------------------------------------------
    # DOCUMENT CONTENT GENERATION
    # -------------------------------------------------------------

    # Main Title
    add_main_title("Smart College Analytics System (SCAS)")
    add_subtitle("Project Profile, System Study, Analysis & Development Environment")

    add_callout(
        "This document contains comprehensive information regarding the Smart College Analytics System (SCAS / MCA-PS3), structured into Introduction, System Study & Analysis, and Development Environment.",
        "PROJECT DOCUMENTATION"
    )

    # =============================================================
    # SECTION 1: INTRODUCTION
    # =============================================================
    add_h1("1 INTRODUCTION")
    add_p("The Smart College Analytics System (SCAS) is an enterprise-grade web application and continuous assessment analytics portal developed for higher education institutions. It connects Students, Faculty Members, and System Administrators into a unified digital ecosystem to streamline academic tracking, attendance management, internal mark processing, fee collection, and career placement preparation.")
    add_p("Goal of the System: To replace manual, paper-based academic workflows with an automated, cloud-ready digital platform that improves operational speed, data accuracy, performance insights, and overall institutional efficiency.")

    # 1.1 Project Profile
    add_h2("1.1 Project Profile")
    add_bullet("Smart College Analytics System (SCAS)", "Project Title: ")
    add_bullet("Web-Based Education ERP & Academic Analytics Portal", "Project Category: ")
    add_bullet("Master of Computer Applications (MCA - Practical System 3)", "Academic Program: ")
    add_bullet("React 19, TypeScript, Vite, Tailwind CSS, Chart.js, Lucide Icons", "Front-End Environment: ")
    add_bullet("Python 3.x, Flask WSGI Framework, Flask-SQLAlchemy (ORM)", "Back-End Environment: ")
    add_bullet("MySQL Server / SQLite Relational Database Engine", "Database System: ")
    add_bullet("Python smtplib, MIME Transport, Local Sent Email Logger", "Email & Notification Engine: ")
    add_bullet("Administrator, Faculty, and Enrolled Students", "Target User Roles: ")

    # 1.2 About the Organization
    add_h2("1.2 About the Organization")
    add_bullet(" The project is developed for the Department of Computer Science & Applications, dedicated to providing high-quality computer applications education, research excellence, and continuous evaluation.", "Academic Institution: ")
    add_bullet(" Transitioning from legacy manual paper registers and spreadsheet-based mark calculations to a modern, centralized cloud-ready digital infrastructure.", "Digital Transformation Goal: ")
    add_bullet(" Empowering students and faculty with real-time academic performance dashboards, interactive SGPA/CGPA analytics graphs, and attendance risk alerts.", "Student-Centric Analytics: ")
    add_bullet(" Enhancing student career readiness through integrated mock placement aptitude tests, digital fee payment tracking, and transparent leave application workflows.", "Career & Operations Focus: ")

    # =============================================================
    # SECTION 2: SYSTEM STUDY AND ANALYSIS
    # =============================================================
    add_h1("2 SYSTEM STUDY AND ANALYSIS")

    # 2.1 Problem Statement
    add_h2("2.1 Problem Statement")
    add_p("Modern educational institutions struggle with fragmented, paper-based administrative operations. Manual handling of student attendance, internal marks processing, fee tracking, and leave requests leads to significant operational bottlenecks:")
    add_bullet(" Manual entry in attendance registers and Excel sheets frequently causes calculation mistakes in end-semester eligibility percentages.", "High Calculation Error Rates: ")
    add_bullet(" Faculty and administrators lack visibility into low-attendance students (<75%) until the end of the term, preventing timely interventions.", "Absence of Real-Time Attendance Warnings: ")
    add_bullet(" Results and internal test marks are calculated manually, leading to long delays in grade publishing and SGPA/CGPA aggregation.", "Delayed Mark Processing: ")
    add_bullet(" Student fee management relies on physical ledgers and offline counters, causing long student queues and payment status tracking errors.", "Manual Fee Record-Keeping: ")
    add_bullet(" Written paper leave and duty applications get misplaced easily, leading to delayed faculty/admin approvals and unrecorded student absences.", "Paper-Based Leave Requests: ")

    # 2.2 Existing System
    add_h2("2.2 Existing System")
    add_p("The existing academic workflow in traditional colleges relies heavily on manual methods, physical registers, isolated Microsoft Excel spreadsheets, paper notices on bulletin boards, and offline fee payment counters.")

    # 2.2.1 Drawbacks
    add_h3("2.2.1 Drawbacks of the Existing System")
    add_bullet(" Student records, attendance details, and mark sheets are stored separately across departments, causing data duplication and conflicting entries.", "1. Data Redundancy & Inconsistency: ")
    add_bullet(" Faculty spend significant manual hours compiling register entries, calculating SGPA/CGPA scores, and checking attendance eligibility.", "2. High Administrative Overhead: ")
    add_bullet(" No graphical tools or interactive charts exist to analyze student performance trends or identify struggling students early.", "3. Lack of Visual Analytics: ")
    add_bullet(" Important college announcements, timetable changes, and fee due reminders rely on physical flyers, which fail to reach students promptly.", "4. Delayed Communication: ")
    add_bullet(" Traditional academic systems operate in isolation from career recruitment activities, lacking built-in tools for placement preparation.", "5. No Placement Test Integration: ")

    # 2.3 Proposed System
    add_h2("2.3 Proposed System")
    add_p("The proposed Smart College Analytics System (SCAS) introduces a modern, web-based platform with role-segmented portals for Students, Faculty, and Administrators. It automates attendance logging, internal mark processing, fee receipt generation, notice broadcasting, and placement readiness into a unified digital dashboard.")

    # 2.3.1 Advantages / Comparison Study
    add_h3("2.3.1 Advantages / Comparison Study")
    add_p("Below is a detailed comparison highlighting the key advantages of the Proposed SCAS Platform over the Existing Manual System:")

    headers_comp = ["Aspect / Feature", "Existing Manual System", "Proposed SCAS Platform"]
    comp_data = [
        ["Attendance Tracking", "Paper registers, manual tallying at end of semester", "Real-time session marking with automated low-attendance warnings (<75%)"],
        ["Mark & GPA Processing", "Manual Excel sheets, error-prone grade calculation", "Automatic SGPA/CGPA computation with dynamic Chart.js analytics graphs"],
        ["Notice Board & Alerts", "Physical paper flyers pinned on college walls", "Digital Notice Board with instant pinned notices and SMTP email alerts"],
        ["Leave & Duty Requests", "Physical written letters submitted manually", "Multi-tier online leave application submission and real-time faculty approval"],
        ["Fee & Payment Dues", "Manual ledger books and separate offline counter", "Real-time fee status tracking (Paid/Pending) with auto-generated receipt numbers"],
        ["Placement Preparation", "No integrated system for placement preparation", "Built-in mock placement aptitude test engine with instant scoreboards"],
        ["User Access & Security", "Unsecured physical records accessible to anyone", "Role-Based Access Control (RBAC) with hashed passwords and session protection"]
    ]
    t_comp = doc.add_table(rows=1, cols=3)
    style_table(t_comp, [2.0, 2.2, 2.3], headers_comp, comp_data)

    # 2.3.2 Modules (with short description)
    add_h3("2.3.2 System Modules (with short description)")

    modules = [
        ("1. Authentication & Security Module", "Enforces Role-Based Access Control (RBAC) with password hashing (Werkzeug/BCrypt) and secure session management for Admin, Faculty, and Student portals."),
        ("2. Administrator Management Module", "Central executive dashboard for registering new students and faculty, dispatching automated credentials, publishing notices, and tracking fee collections."),
        ("3. Faculty Management Module", "Manages subject assignments, real-time daily student attendance marking, internal/midterm marks entry, placement exam creation, and leave approvals."),
        ("4. Student Portal Module", "Provides students with an academic overview, SGPA/CGPA result sheets, subject-wise attendance tracking, visual performance analytics, and leave application tools."),
        ("5. Fee & Payment Management Module", "Tracks student tuition fee records, updates payment statuses ('Paid' or 'Pending'), and automatically generates unique transaction receipt numbers."),
        ("6. Placement & Aptitude Test Module", "Features an interactive mock exam engine where faculty can design multiple-choice tests and students can practice for campus placement drives."),
        ("7. Email & Notification Engine", "Integrates SMTP mail services to dispatch welcome credentials, temporary password resets, and high-priority institutional announcements.")
    ]

    for title, desc in modules:
        add_bullet(desc, f"{title}: ")

    # =============================================================
    # SECTION 3: DEVELOPMENT ENVIRONMENT
    # =============================================================
    add_h1("3 DEVELOPMENT ENVIRONMENT")

    # 3.1 Hardware Requirement
    add_h2("3.1 Hardware Requirement")
    add_p("To run the SCAS application smoothly across development, testing, and client deployment environments, the following hardware specifications are required:")

    add_bullet("Intel Core i3 / i5 / AMD Ryzen 5 processor (2.0 GHz or higher).", "Processor (CPU): ")
    add_bullet("Minimum 4 GB (8 GB recommended for simultaneous backend and frontend servers).", "System Memory (RAM): ")
    add_bullet("Minimum 20 GB available storage space (SSD recommended for fast build performance).", "Storage (HDD/SSD): ")
    add_bullet("1366 x 768 minimum resolution (1920 x 1080 Full HD recommended).", "Display Monitor: ")
    add_bullet("Standard LAN / Wi-Fi internet connection for network access and SMTP transport.", "Network Interface: ")

    # 3.2 Software Requirement
    add_h2("3.2 Software Requirement")
    add_p("The software environment required to develop, build, and deploy the SCAS platform includes:")

    add_bullet("Windows 10 / 11, Linux (Ubuntu/Debian), or macOS.", "Operating System: ")
    add_bullet("Python 3.10+ (for Flask backend) and Node.js v18+ (for Vite frontend).", "Runtime Environments: ")
    add_bullet("Flask WSGI Server (Development) / Gunicorn WSGI Server (Production).", "Web Server: ")
    add_bullet("MySQL Server 8.0 (via XAMPP / WAMP) or SQLite3 Relational Database Engine.", "Database Server: ")
    add_bullet("Visual Studio Code (VS Code), Git Version Control, Postman, Vite CLI Tooling.", "Development Tools / IDE: ")
    add_bullet("Google Chrome, Mozilla Firefox, or Microsoft Edge.", "Web Browser: ")

    # 3.3 Programming Environment
    add_h2("3.3 Programming Environment")

    # 3.3.1 About Front-End
    add_h3("3.3.1 About Front-End")
    add_p("The front-end of SCAS is designed as a modern, responsive Single Page Application (SPA) providing an interactive user experience across desktop and mobile browsers:")

    add_bullet("Component-driven architecture providing strict type safety, modular code organization, and fast state synchronization.", "React 19 & TypeScript: ")
    add_bullet("Next-generation frontend tooling offering superfast Hot Module Replacement (HMR) during development and optimized production bundling.", "Vite Build Engine: ")
    add_bullet("Utility-first CSS framework enabling sleek dark mode glassmorphism, responsive grid layouts, and custom micro-animations.", "Tailwind CSS: ")
    add_bullet("Renders dynamic visual charts for GPA progression trends, subject-wise attendance distribution, and performance radar charts.", "Chart.js & react-chartjs-2: ")
    add_bullet("Provides clean, intuitive vector iconography across all dashboard panels and navigation menus.", "Lucide React Icons: ")

    # 3.3.2 About Back-End
    add_h3("3.3.2 About Back-End")
    add_p("The back-end of SCAS is engineered with Python and Flask, delivering a RESTful web service architecture for secure data handling and business logic execution:")

    add_bullet("Lightweight, extensible Python WSGI framework managing RESTful API routes and blueprint controllers (`auth`, `admin`, `faculty`, `student`).", "Python 3.x & Flask Framework: ")
    add_bullet("Object-Relational Mapping (ORM) library translating Python model classes into SQL queries for database operations.", "Flask-SQLAlchemy (ORM): ")
    add_bullet("Relational database storing 12 normalized tables with strict primary keys, foreign key constraints, and cascade delete rules.", "MySQL Database Engine: ")
    add_bullet("Provides secure password hashing (`generate_password_hash`) and checking for safe user authentication.", "Werkzeug Security: ")
    add_bullet("Automated email dispatch module for sending account activation credentials and password reset notifications via SMTP with local logging fallback.", "SMTP Mail Transport: ")

    # Save Document
    doc_path = os.path.join(os.getcwd(), "info.docx")
    doc.save(doc_path)
    print(f"Info document successfully created at: {doc_path}")

if __name__ == "__main__":
    create_info_document()
